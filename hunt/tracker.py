"""Excel tracker (schema in tracker/README.md). Dedup key: URL, including the 'Autres URLs' column."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .models import Listing

SHEET = "Annonces"
COLUMNS = [
    "URL", "Plateforme", "Date d'ajout", "Prix", "Surface (m²)", "Pièces", "Chambres", "Code postal",
    "Ville / Quartier", "Métro", "Ligne 14 (m)", "Étage", "Ascenseur", "Meublé", "Type de contact", "Disponible le", "Zone",
    "Priorité", "Statut", "Message envoyé", "Autres URLs", "Notes",
]
FILLS = {"HAUTE": "E2EFDA", "MOYENNE": "FFFFC7", "BASSE": "FCE4D6"}


def _open(path: Path):
    if path.exists():
        wb = load_workbook(path)
        ws = wb[SHEET]
        head = [c.value for c in ws[1]]
        for name in COLUMNS:  # trackers created before a column existed get it appended at the end
            if name not in head:
                ws.cell(row=1, column=len(head) + 1, value=name).font = Font(bold=True)
                head.append(name)
        return wb, ws
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    return wb, ws


def known_urls(path: Path) -> set[str]:
    if not path.exists():
        return set()
    ws = load_workbook(path, read_only=True)[SHEET]
    head = [c.value for c in next(ws.iter_rows(max_row=1))]
    iu, io = head.index("URL"), head.index("Autres URLs") if "Autres URLs" in head else None
    urls = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[iu]:
            urls.add(row[iu])
        if io is not None and row[io]:
            urls.update(u.strip() for u in str(row[io]).split("\n") if u.strip())
    return urls


def append(path: Path, listings: list[Listing]) -> None:
    wb, ws = _open(path)
    today = date.today().isoformat()
    yn = {True: "oui", False: "non", None: ""}
    head = [c.value for c in ws[1]]
    for l in listings:
        metro = f"{l.metro_station} ({', '.join(l.metro_lines or [])}) — {l.metro_m} m" if l.metro_station else None
        values = dict(zip(COLUMNS, [
            l.url, l.source, today, l.price, l.surface, l.rooms, l.bedrooms, l.postal_code,
            " / ".join(x for x in (l.city, l.district) if x), metro, l.line14_m, l.floor, yn[l.elevator], yn[l.furnished],
            l.contact_type, l.available_from, l.zone_tier, None, "Nouveau", "non", "\n".join(l.other_urls), None,
        ]))
        ws.append([values.get(h) for h in head])
    wb.save(path)


def set_priorities(path: Path, priorities: dict[str, str]) -> int:
    """priorities: {url: 'HAUTE'|'MOYENNE'|'BASSE'}. Writes the column and colours the row. Returns rows updated."""
    wb, ws = _open(path)
    head = [c.value for c in ws[1]]
    iu, ip = head.index("URL"), head.index("Priorité")
    n = 0
    for row in ws.iter_rows(min_row=2):
        p = priorities.get(row[iu].value)
        if p:
            row[ip].value = p
            fill = PatternFill("solid", fgColor=FILLS.get(p, "FFFFFF"))
            for c in row:
                c.fill = fill
            n += 1
    wb.save(path)
    return n
